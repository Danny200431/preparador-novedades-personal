
# Este archivo crea el reporte de Excel con los resultados
# de las novedades de personal.

from pathlib import Path
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# Ubicación principal del proyecto.
CARPETA_PROYECTO = Path(__file__).resolve().parent.parent


def clasificar_rechazo(motivo):
    """Identifica si el rechazo es de validación o de negocio."""

    texto = str(motivo or "").strip().lower()

    motivos_regla_negocio = (
        "el área del cargo no coincide con el centro de costo",
        "el area del cargo no coincide con el centro de costo",
        "el cargo no existe en el catálogo",
        "el cargo no existe en el catalogo",
        "el centro de costo no existe en el catálogo",
        "el centro de costo no existe en el catalogo",
        "el cargo y el área no están permitidos",
        "el cargo y el area no estan permitidos",
    )

    if texto in motivos_regla_negocio:
        return "REGLA_NEGOCIO"

    return "VALIDACION"


def obtener_valor_recibido(resultado, datos, motivo, campo):
    """Busca el dato que causó el rechazo."""

    novedad_original = resultado.get("novedad_original", {})

    if not isinstance(novedad_original, dict):
        novedad_original = {}

    fuentes = [novedad_original, resultado, datos]

    for fuente in fuentes:
        valor = fuente.get("valor_recibido")

        if valor is not None and str(valor).strip() != "":
            return valor

    texto = str(motivo or "").strip().lower()
    campo = str(campo or "").strip().lower()

    if "falta el campo obligatorio:" in texto:
        nombre = texto.split(
            "falta el campo obligatorio:", 1
        )[1].strip()

        for fuente in fuentes:
            valor = fuente.get(nombre)

            if valor is not None and str(valor).strip() != "":
                return valor

        return "Campo vacío"

    if "la cc debe tener" in texto:
        campos = ["numero_documento"]
    elif any(frase in texto for frase in (
        "fecha efectiva ya pasó",
        "fecha efectiva ya paso",
        "fecha supera el límite permitido",
        "fecha supera el limite permitido",
    )):
        campos = ["fecha_efectiva"]
    elif (
        "el área del cargo no coincide" in texto
        or "el area del cargo no coincide" in texto
    ):
        campos = ["centro_costo", "cargo"]
    elif "el centro de costo no existe" in texto:
        campos = ["centro_costo"]
    elif "el cargo no existe" in texto:
        campos = ["cargo", "codigo_cargo"]
    elif "teléfono" in texto or "telefono" in texto:
        campos = ["telefono"]
    else:
        equivalencias = {
            "tipo": ["tipo", "tipo_novedad"],
            "tipo_novedad": ["tipo_novedad", "tipo"],
            "numero_documento": ["numero_documento"],
            "documento": ["numero_documento"],
            "fecha_efectiva": ["fecha_efectiva"],
            "centro_costo": ["centro_costo"],
            "cargo": ["cargo", "codigo_cargo"],
            "cargo/centro_costo": [
                "cargo", "codigo_cargo", "centro_costo"
            ],
            "telefono": ["telefono"],
            "correo_corporativo": ["correo_corporativo"],
            "nombre": ["nombre_completo", "nombre", "nombres"],
        }

        campos = equivalencias.get(campo, [])

    for nombre in campos:
        for fuente in fuentes:
            valor = fuente.get(nombre)

            if valor is not None and str(valor).strip() != "":
                return valor

    return ""


def agregar_hoja(libro, nombre, columnas, filas):
    """Crea una hoja y le aplica formato básico."""

    hoja = libro.create_sheet(title=nombre)
    hoja.append(columnas)

    for fila in filas:
        hoja.append([fila.get(columna, "") for columna in columnas])

    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions

    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78"
        )
        celda.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    for columna in hoja.columns:
        letra = get_column_letter(columna[0].column)
        ancho = max(
            len(str(celda.value or ""))
            for celda in columna
        )
        hoja.column_dimensions[letra].width = min(ancho + 2, 45)

    return hoja


def crear_reporte(resultados, parametros, archivo_entrada=""):
    """Crea el archivo Excel con las cuatro hojas del reporte."""

    carpeta_salida = CARPETA_PROYECTO / "salida"
    carpeta_salida.mkdir(parents=True, exist_ok=True)

    aceptadas = []
    rechazadas = []
    correos = {}

    # Separamos las novedades aceptadas y rechazadas.
    for resultado in resultados:
        datos = resultado.get("datos", {})

        if not isinstance(datos, dict):
            datos = {}

        if resultado.get("resultado") == "ACEPTADA":
            aceptadas.append(resultado)
            continue

        motivo = resultado.get(
            "motivo", resultado.get("mensaje", "")
        )

        clasificacion = str(
            resultado.get(
                "clasificacion",
                datos.get("clasificacion", "")
            )
        ).strip().upper()

        if clasificacion not in ("VALIDACION", "REGLA_NEGOCIO"):
            clasificacion = clasificar_rechazo(motivo)

        campo = resultado.get("campo", datos.get("campo", ""))

        rechazadas.append({
            "consecutivo": resultado.get("consecutivo", ""),
            "tipo": resultado.get(
                "tipo", resultado.get("tipo_novedad", "")
            ),
            "numero_documento": resultado.get(
                "numero_documento", ""
            ),
            "nombre": resultado.get(
                "nombre_completo", resultado.get("nombre", "")
            ),
            "clasificacion": clasificacion,
            "campo": campo,
            "motivo": motivo,
            "valor_recibido": obtener_valor_recibido(
                resultado, datos, motivo, campo
            ),
        })

    # Preparamos las columnas de las novedades aceptadas.
    columnas_carga = [
        "consecutivo", "tipo_novedad", "tipo_documento",
        "numero_documento", "nombre_completo",
        "correo_corporativo", "codigo_cargo", "nombre_cargo",
        "nivel", "centro_costo", "area", "perfil_sistema",
        "requiere_aprobacion", "fecha_efectiva", "telefono",
        "fecha_vinculacion_previa",
    ]

    filas_carga = []

    for resultado in aceptadas:
        datos = resultado.get("datos", {})

        if not isinstance(datos, dict):
            datos = {}

        fila = {}

        for columna in columnas_carga:
            if columna == "consecutivo":
                fila[columna] = resultado.get("consecutivo", "")
            else:
                fila[columna] = datos.get(
                    columna, resultado.get(columna, "")
                )

        filas_carga.append(fila)

    # Revisamos si hay correos corporativos repetidos.
    for resultado in aceptadas:
        datos = resultado.get("datos", {})

        if not isinstance(datos, dict):
            datos = {}

        correo = str(
            datos.get(
                "correo_corporativo",
                resultado.get("correo_corporativo", "")
            )
        ).strip().lower()

        if correo and "@" in correo:
            base = correo.split("@", 1)[0]
            correos.setdefault(base, []).append({
                "consecutivo": resultado.get("consecutivo", ""),
                "correo": correo,
            })

    filas_colisiones = []

    for base, registros in correos.items():
        if len(registros) > 1:
            filas_colisiones.append({
                "correo_base": base,
                "consecutivos": ", ".join(
                    str(registro["consecutivo"])
                    for registro in registros
                ),
                "correos_finales": ", ".join(
                    registro["correo"] for registro in registros
                ),
            })

    # Contamos las novedades por tipo.
    tipos = {}

    for resultado in resultados:
        tipo = str(resultado.get("tipo", "Sin tipo")).strip().upper()

        if tipo not in tipos:
            tipos[tipo] = {
                "tipo_novedad": tipo,
                "recibidas": 0,
                "validas": 0,
                "rechazadas": 0,
            }

        tipos[tipo]["recibidas"] += 1

        if resultado.get("resultado") == "ACEPTADA":
            tipos[tipo]["validas"] += 1
        else:
            tipos[tipo]["rechazadas"] += 1

    filas_resumen = []

    for tipo, datos_tipo in tipos.items():
        recibidas = datos_tipo["recibidas"]
        validas = datos_tipo["validas"]

        filas_resumen.append({
            "tipo_novedad": tipo,
            "recibidas": recibidas,
            "validas": validas,
            "rechazadas": datos_tipo["rechazadas"],
            "porcentaje_exito": (
                round(validas / recibidas * 100, 2)
                if recibidas else 0
            ),
        })

    filas_resumen.append({
        "tipo_novedad": "TOTAL",
        "recibidas": len(resultados),
        "validas": len(aceptadas),
        "rechazadas": len(rechazadas),
        "porcentaje_exito": (
            round(len(aceptadas) / len(resultados) * 100, 2)
            if resultados else 0
        ),
    })

    motivos = [
        registro["motivo"]
        for registro in rechazadas
        if registro["motivo"]
    ]

    motivo_frecuente = ""

    if motivos:
        cantidades = {}

        for motivo in motivos:
            cantidades[motivo] = cantidades.get(motivo, 0) + 1

        motivo_frecuente = max(cantidades, key=cantidades.get)

    filas_resumen.append({
        "tipo_novedad": "Motivo de rechazo más frecuente",
        "recibidas": "",
        "validas": "",
        "rechazadas": "",
        "porcentaje_exito": motivo_frecuente,
    })

    filas_resumen.append({
        "tipo_novedad": "Archivo de entrada",
        "recibidas": "",
        "validas": "",
        "rechazadas": "",
        "porcentaje_exito": Path(archivo_entrada).name
        if archivo_entrada else "",
    })

    # Usamos el nombre del parámetro que existe en el Excel.
    prefijo = parametros.get("archivo_salida")

    if not prefijo:
        raise ValueError(
            "Falta el parámetro 'archivo_salida' "
            "en la hoja Parametros del Excel de configuración."
        )

    fecha_actual = datetime.now().strftime("%Y%m%d")
    ruta_salida = carpeta_salida / f"{prefijo}_{fecha_actual}.xlsx"

    # Creamos el Excel directamente con openpyxl.
    libro = Workbook()
    libro.remove(libro.active)

    agregar_hoja(
        libro,
        "Carga_Sistema",
        columnas_carga,
        filas_carga
    )

    agregar_hoja(
        libro,
        "Rechazados",
        [
            "consecutivo", "tipo", "numero_documento", "nombre",
            "clasificacion", "campo", "motivo", "valor_recibido",
        ],
        rechazadas
    )

    agregar_hoja(
        libro,
        "Colisiones_Correo",
        ["correo_base", "consecutivos", "correos_finales"],
        filas_colisiones
    )

    agregar_hoja(
        libro,
        "Resumen",
        [
            "tipo_novedad", "recibidas", "validas",
            "rechazadas", "porcentaje_exito",
        ],
        filas_resumen
    )

    libro.save(ruta_salida)
    libro.close()

    return str(ruta_salida)
