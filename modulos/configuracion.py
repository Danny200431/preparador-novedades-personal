
# Este archivo lee la información del Excel de configuración.

from pathlib import Path
from openpyxl import load_workbook
from modulos.excepciones import ErrorProyecto


# Buscamos la carpeta principal del proyecto.
CARPETA_PROYECTO = Path(__file__).resolve().parent.parent


def cargar_parametros():
    ruta = CARPETA_PROYECTO / "config" / "Parametros.xlsx"

    if not ruta.is_file():
        raise ErrorProyecto(
            "No se encontró el archivo config/Parametros.xlsx."
        )

    excel = None

    try:
        excel = load_workbook(ruta, data_only=True)

        hojas_necesarias = [
            "Parametros",
            "Cargos",
            "Centros_Costo",
            "Matriz_Roles",
            "Campos_Obligatorios"
        ]

        for nombre in hojas_necesarias:
            if nombre not in excel.sheetnames:
                raise ErrorProyecto(
                    f"Falta la hoja '{nombre}' en el Excel de configuración."
                )

        # Leemos los parámetros generales.
        hoja = excel["Parametros"]
        parametros = {}

        for fila in hoja.iter_rows(min_row=2, values_only=True):
            if fila[0] is not None:
                clave = (
                    str(fila[0])
                    .replace("\xa0", " ")
                    .strip()
                    .lower()
                )
                valor = fila[1] if len(fila) > 1 else None
                parametros[clave] = valor

        # Comprobamos el nombre correcto del archivo de salida.
        prefijo = parametros.get("archivo_salida")

        if prefijo is None or not str(prefijo).strip():
            raise ErrorProyecto(
                "Falta el parámetro 'archivo_salida' en la hoja Parametros."
            )

        # Leemos los cargos.
        hoja = excel["Cargos"]
        cargos = {}

        for fila in hoja.iter_rows(min_row=2, values_only=True):
            if fila[0] is not None:
                codigo = str(fila[0]).strip().upper()
                cargos[codigo] = fila[1:]

        # Leemos los centros de costo.
        hoja = excel["Centros_Costo"]
        centros = {}

        for fila in hoja.iter_rows(min_row=2, values_only=True):
            if fila[0] is not None:
                codigo = str(fila[0]).strip().upper()
                centros[codigo] = fila[1:]

        # Leemos la matriz de roles.
        hoja = excel["Matriz_Roles"]
        matriz = []

        for fila in hoja.iter_rows(min_row=2, values_only=True):
            if fila[0] is not None:
                fila_limpia = tuple(
                    valor.strip() if isinstance(valor, str) else valor
                    for valor in fila
                )
                matriz.append(fila_limpia)

        # Leemos los campos obligatorios.
        hoja = excel["Campos_Obligatorios"]
        campos_obligatorios = {}

        for fila in hoja.iter_rows(min_row=2, values_only=True):
            if fila[0] is not None and fila[1] is not None:
                tipo = str(fila[0]).strip().upper()
                campos = str(fila[1]).split(";")
                campos_obligatorios[tipo] = [
                    campo.strip() for campo in campos
                ]

        return (
            parametros,
            cargos,
            centros,
            matriz,
            campos_obligatorios
        )

    finally:
        if excel is not None:
            excel.close()
