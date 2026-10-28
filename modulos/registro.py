# Importamos las herramientas necesarias
import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


# Guardamos los mensajes de la ejecución actual
mensajes = []

# Definimos los niveles de importancia de los mensajes
NIVELES = {
    "INFO": 10,
    "WARNING": 20,
    "ERROR": 30,
    "CRITICAL": 40
}


def guardar_mensaje(
    mensaje,
    nivel="INFO",
    modulo="",
    funcion="",
    registro="",
    detalle_tecnico=""
):
    # Guardamos la información de cada evento
    mensajes.append({
        "marca_tiempo": datetime.now().strftime(
            "%Y-%m-%d %H:%M:%S"
        ),
        "nivel": nivel.upper(),
        "modulo": modulo,
        "funcion": funcion,
        "registro": registro,
        "mensaje": str(mensaje),
        "detalle_tecnico": str(detalle_tecnico)
    })


def mostrar_mensajes():
    # Mostramos los mensajes guardados en la terminal
    for mensaje in mensajes:
        print(
            mensaje["nivel"],
            "-",
            mensaje["mensaje"]
        )


def crear_registro(
    resultados,
    parametros,
    inicio=None,
    archivo_entrada=""
):
    # Guardamos la hora de finalización
    fin = datetime.now()

    # Si no recibimos la hora de inicio, usamos la hora actual
    if inicio is None:
        inicio = fin

    # Calculamos la duración de la ejecución en segundos
    duracion = round((fin - inicio).total_seconds(), 2)

    # Obtenemos el prefijo del nombre del archivo desde la configuración
    prefijo = parametros.get(
        "prefijo_log",
        "log_ejecucion"
    )

    # Obtenemos el nivel mínimo de los mensajes que se guardarán
    nivel_minimo = str(
        parametros.get("nivel_log_minimo", "INFO")
    ).upper()

    # Si el nivel configurado no existe, utilizamos INFO
    if nivel_minimo not in NIVELES:
        nivel_minimo = "INFO"

    # Creamos un nombre con fecha y hora para el archivo
    nombre = (
        f"{prefijo}_"
        f"{fin.strftime('%Y%m%d_%H%M%S')}.xlsx"
    )

    # Creamos la carpeta de registros si todavía no existe
    os.makedirs("logs", exist_ok=True)

    # Construimos la ruta donde se guardará el archivo
    ruta = os.path.join("logs", nombre)

    # Creamos el libro de Excel
    libro = Workbook()

    # Creamos la hoja Detalle
    detalle = libro.active
    detalle.title = "Detalle"

    # Escribimos los encabezados de la hoja Detalle
    detalle.append([
        "Consecutivo",
        "Marca de tiempo",
        "Nivel",
        "Modulo",
        "Funcion",
        "Registro",
        "Mensaje",
        "Detalle_tecnico"
    ])

    # Iniciamos la numeración de los eventos en 1
    consecutivo_log = 1

    # Guardamos los mensajes generados durante la ejecución
    for mensaje in mensajes:
        nivel = mensaje.get("nivel", "INFO")

        # Comprobamos si el mensaje cumple el nivel mínimo configurado
        if (
            nivel in NIVELES
            and NIVELES[nivel] >= NIVELES[nivel_minimo]
        ):
            detalle.append([
                consecutivo_log,
                mensaje.get("marca_tiempo", ""),
                nivel,
                mensaje.get("modulo", ""),
                mensaje.get("funcion", ""),
                mensaje.get("registro", ""),
                mensaje.get("mensaje", ""),
                mensaje.get("detalle_tecnico", "")
            ])

            # Aumentamos el número para el siguiente evento
            consecutivo_log += 1

    # Contamos cuántas novedades fueron aceptadas
    aceptadas = sum(
        1 for resultado in resultados
        if resultado.get("resultado") == "ACEPTADA"
    )

    # Contamos cuántas novedades fueron rechazadas
    rechazadas = sum(
        1 for resultado in resultados
        if resultado.get("resultado") == "RECHAZADA"
    )

    # Registramos el resultado de cada novedad
    for resultado in resultados:
        estado = resultado.get("resultado", "")

        # Las aceptadas se registran como INFO y las rechazadas como WARNING
        nivel = "INFO" if estado == "ACEPTADA" else "WARNING"

        # Solo guardamos los resultados que cumplen el nivel mínimo
        if NIVELES[nivel] >= NIVELES[nivel_minimo]:
            detalle.append([
                consecutivo_log,
                fin.strftime("%Y-%m-%d %H:%M:%S"),
                nivel,
                "procesamiento",
                "procesar_novedad",
                resultado.get("numero_documento", ""),
                resultado.get("mensaje", ""),
                ""
            ])

            # Continuamos la numeración sin repetir ni saltar números
            consecutivo_log += 1

    # Creamos la hoja Resumen
    resumen = libro.create_sheet("Resumen")

    # Escribimos los datos generales de la ejecución
    resumen.append(["Concepto", "Valor"])

    resumen.append([
        "Inicio",
        inicio.strftime("%Y-%m-%d %H:%M:%S")
    ])

    resumen.append([
        "Fin",
        fin.strftime("%Y-%m-%d %H:%M:%S")
    ])

    resumen.append([
        "Duración en segundos",
        duracion
    ])

    resumen.append([
        "Archivo de entrada",
        archivo_entrada
    ])

    # Contamos los eventos de cada nivel
    resumen.append([
        "Eventos INFO",
        sum(
            1 for mensaje in mensajes
            if mensaje.get("nivel") == "INFO"
        )
    ])

    resumen.append([
        "Eventos WARNING",
        sum(
            1 for mensaje in mensajes
            if mensaje.get("nivel") == "WARNING"
        )
    ])

    resumen.append([
        "Eventos ERROR",
        sum(
            1 for mensaje in mensajes
            if mensaje.get("nivel") == "ERROR"
        )
    ])

    resumen.append([
        "Eventos CRITICAL",
        sum(
            1 for mensaje in mensajes
            if mensaje.get("nivel") == "CRITICAL"
        )
    ])

    # Registramos el total de novedades recibidas, válidas y rechazadas
    resumen.append([
        "Total recibido",
        len(resultados)
    ])

    resumen.append([
        "Total válido",
        aceptadas
    ])

    resumen.append([
        "Total rechazado",
        rechazadas
    ])

    # Damos formato a las hojas Detalle y Resumen
    for hoja in [detalle, resumen]:

        # Aplicamos color y negrita a los encabezados
        for celda in hoja[1]:
            celda.font = Font(
                bold=True,
                color="FFFFFF"
            )

            celda.fill = PatternFill(
                fill_type="solid",
                fgColor="4472C4"
            )

        # Ajustamos el ancho de las columnas según su contenido
        for columna in hoja.columns:
            letra = columna[0].column_letter

            ancho = max(
                len(str(celda.value or ""))
                for celda in columna
            )

            hoja.column_dimensions[letra].width = min(
                ancho + 3,
                45
            )

        # Dejamos fija la fila de encabezados al desplazarnos
        hoja.freeze_panes = "A2"

    # Pintamos de rojo las filas que tengan errores graves
    for fila in detalle.iter_rows(min_row=2):
        nivel = fila[2].value

        if nivel in ("ERROR", "CRITICAL"):
            for celda in fila:
                celda.fill = PatternFill(
                    fill_type="solid",
                    fgColor="FFC7CE"
                )

    # Guardamos el libro de Excel en la carpeta logs
    libro.save(ruta)

    # Devolvemos la ruta del archivo creado
    return ruta